"""Engine tests on small made-up workbooks (no real patient data)."""
import datetime as dt
import io

import openpyxl
import pytest

import vlookup_engine as ve

HEADERS = ["PATIENT NAME", "DOB", "DOS", "CPT", "STATUS", "COMMENTS"]
LOOKUP_ROWS = [
    ["DOE, JANE", dt.datetime(1950, 1, 2), dt.datetime(2026, 1, 5), "99213", "PM", "ok"],
    ["ROE, RICHARD", dt.datetime(1960, 3, 4), dt.datetime(2026, 2, 6), "99214", "PNM", "needs note"],
    ["POE, EDGAR", dt.datetime(1970, 5, 6), dt.datetime(2026, 3, 7), "99215", "PM", ""],
]


def workbook(sheets: dict[str, list[list]]) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def target_rows(extra=()):
    """Same patients without STATUS / COMMENTS, DOB as text, plus any extra rows."""
    rows = [[r[0].lower(), r[1].strftime("%m/%d/%Y"), r[2], r[3]] for r in LOOKUP_ROWS]
    return [HEADERS[:4]] + rows + list(extra)


def run(target_bytes, lookup_bytes, opts=None, sheet="001", **job_kwargs):
    opts = opts or ve.Options()
    target, lookup = ve.read_workbook(target_bytes)[sheet], ve.read_workbook(lookup_bytes)[sheet]
    job = ve.SheetJob(sheet, sheet, job_kwargs.pop("keys", [("PATIENT NAME", "PATIENT NAME"), ("DOB", "DOB")]),
                      job_kwargs.pop("ret", ["STATUS", "COMMENTS"]), **job_kwargs)
    res = ve.run_job(target, lookup, job, opts)
    out = openpyxl.load_workbook(io.BytesIO(ve.build_output(target_bytes, [res], opts)))
    return res, out


@pytest.fixture
def lookup_bytes():
    return workbook({"001": [HEADERS] + LOOKUP_ROWS})


def test_fills_new_columns_matching_case_and_date_formats(lookup_bytes):
    res, out = run(workbook({"001": target_rows()}), lookup_bytes)
    ws = out["001"]
    assert [c.value for c in ws[1]] == HEADERS
    assert [ws.cell(r, 5).value for r in (2, 3, 4)] == ["PM", "PNM", "PM"]
    assert res.counts[ve.MATCHED] == 3 and res.counts_match


def test_not_found_rows_are_blank_and_ignored_by_count_check(lookup_bytes):
    new = ["NEW, PATIENT", "01/01/1990", dt.datetime(2026, 4, 1), "99213"]
    res, out = run(workbook({"001": target_rows([new])}), lookup_bytes)
    assert res.counts[ve.NOT_FOUND] == 1
    assert out["001"].cell(5, 5).value is None
    assert out["001"].cell(5, 5).fill.fgColor.rgb.endswith("FFF4C7")
    assert res.counts_match

    strict, _ = run(workbook({"001": target_rows([new])}), lookup_bytes, ve.Options(count_ignore_blank=False))
    assert not strict.counts_match


def test_count_check_warns_when_lookup_rows_are_missing(lookup_bytes):
    rows = target_rows()[:-1]  # last patient missing from the new file
    res, _ = run(workbook({"001": rows}), lookup_bytes)
    assert res.lookup_unused == 1
    assert not res.counts_match
    assert any("STATUS: PM -1" in p for p in res.count_problems())


def test_existing_na_is_removed_and_values_are_kept(lookup_bytes):
    rows = [HEADERS] + [r[:4] + ["#N/A", "#N/A"] for r in LOOKUP_ROWS]
    rows[2][4] = "KEEP ME"
    res, out = run(workbook({"001": rows}), lookup_bytes)
    values = [c.value for row in out["001"].iter_rows(min_row=2) for c in row]
    assert "#N/A" not in values
    assert out["001"].cell(3, 5).value == "KEEP ME" and res.cells_kept == 1


def test_fallback_key(lookup_bytes):
    rows = target_rows()
    rows[1][1] = "12/31/1999"  # wrong DOB: primary key fails, name alone matches
    res, out = run(workbook({"001": rows}), lookup_bytes, fallback_keys=[("PATIENT NAME", "PATIENT NAME")])
    assert res.counts[ve.FALLBACK] == 1
    assert out["001"].cell(2, 5).value == "PM"


@pytest.mark.parametrize("rule, expected", [("first", "PM"), ("last", "PNM"), ("join", "PM | PNM")])
def test_duplicate_rules(rule, expected):
    lookup = workbook({"001": [HEADERS, LOOKUP_ROWS[0], LOOKUP_ROWS[0][:4] + ["PNM", ""]]})
    res, out = run(workbook({"001": target_rows()[:2]}), lookup, ve.Options(duplicate_rule=rule))
    assert out["001"].cell(2, 5).value == expected
    assert res.conflicting_keys == 1


def test_unique_key_in_both_files_and_all_sheets_kept(lookup_bytes):
    target_bytes = workbook({"MAIN": [["X"], [1]], "001": target_rows(), "130": [["Y"], [2]]})
    cols = [("PATIENT NAME", "PATIENT NAME"), ("DOS", "DOS")]
    res, out = run(target_bytes, lookup_bytes, keys=cols, unique_key="UNIQUE KEY", unique_separator="-",
                   unique_columns=cols)
    assert out.sheetnames == ["MAIN", "001", "130"]
    assert out["MAIN"]["A2"].value == 1 and out["130"]["A2"].value == 2
    assert out["001"].cell(1, 7).value == "UNIQUE KEY"
    assert out["001"].cell(2, 7).value == "doe, jane-01/05/2026"
    lookup_out = openpyxl.load_workbook(io.BytesIO(ve.build_lookup_output(lookup_bytes, [res])))
    assert lookup_out["001"].cell(2, 7).value == "DOE, JANE-01/05/2026"


def test_unique_key_name_cannot_overwrite_a_key_column(lookup_bytes):
    with pytest.raises(ValueError, match="UNIQUE KEY column can't be named"):
        run(workbook({"001": target_rows()}), lookup_bytes, unique_key="patient name")


def test_merged_destination_cells_are_unmerged(lookup_bytes):
    wb = openpyxl.load_workbook(io.BytesIO(workbook({"001": [HEADERS] + [r[:4] for r in LOOKUP_ROWS]})))
    wb["001"].merge_cells("E2:F3")
    buf = io.BytesIO()
    wb.save(buf)
    _, out = run(buf.getvalue(), lookup_bytes)
    assert out["001"].cell(3, 5).value == "PNM"


def test_text_is_made_safe_for_excel():
    assert ve._safe_value("a\x00b\x07c") == "abc"
    assert len(ve._safe_value("x" * 40000)) == ve.EXCEL_MAX_TEXT


def test_header_row_beyond_data_gives_empty_sheet(lookup_bytes):
    sheets = ve.read_workbook(lookup_bytes, header_row=40)
    assert sheets["001"].headers == [] and sheets["001"].rows == []


@pytest.mark.parametrize("data, message", [
    (b"", "empty"),
    (b"PATIENT NAME,DOB\nDOE,1/1/1950\n", "not a valid Excel workbook"),
    (b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, "password-protected"),
    (b"PK\x03\x04garbage", "not a valid Excel workbook"),
])
def test_bad_files_give_friendly_errors(data, message):
    with pytest.raises(ve.WorkbookError, match=message):
        ve.read_workbook(data)


def test_suggestions(lookup_bytes):
    target = ve.read_workbook(workbook({"001": target_rows()}))["001"]
    lookup = ve.read_workbook(lookup_bytes)["001"]
    assert ve.suggest_keys(target, lookup) == ["PATIENT NAME", "DOB"]
    assert ve.suggest_return_columns(target, lookup) == ["STATUS", "COMMENTS"]
    assert ve.suggest_lookup_sheet("001", ["MAIN", "measure_001", "130"]) == "measure_001"
