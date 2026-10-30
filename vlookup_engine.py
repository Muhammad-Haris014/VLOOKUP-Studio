"""VLOOKUP engine: reads workbooks, matches rows on chosen key columns and writes looked-up values.

The workbook to fill is never modified in place; `build_output` returns a new file as bytes.
"""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from collections import Counter
from copy import copy
from dataclasses import dataclass, field

import openpyxl
import pandas as pd
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, MergedCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

NOT_FOUND_FILL = PatternFill("solid", fgColor="FFF4C7")
FALLBACK_FILL = PatternFill("solid", fgColor="FFE2CC")
SUMMARY_SHEET = "VLOOKUP SUMMARY"
MATCH_COLUMN = "MATCH STATUS"

MATCHED = "Matched"
FALLBACK = "Matched (fallback)"
NOT_FOUND = "Not found"
BLANK_KEY = "Blank key"

DUPLICATE_RULES = {
    "first": "First match (same as Excel VLOOKUP)",
    "first_nonblank": "First non-blank value",
    "last": "Last match",
    "join": "Join all unique values",
}

BLANK_LABEL = "(blank)"
EXCEL_ERRORS = {"#N/A", "#N/A N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}

EXCEL_MAX_TEXT = 32767  # longest text Excel allows in one cell

_KEEP = object()  # marker: leave the existing cell untouched
_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # old .xls files and password-protected workbooks
_SLASH_DATE = re.compile(r"^(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2}|\d{4})$")
_ISO_DATE = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[ T]00:00(?::00)?)?$")


# --------------------------------------------------------------------------- reading


def header_key(header) -> str:
    return re.sub(r"\s+", " ", str(header)).strip().upper()


@dataclass
class Sheet:
    name: str
    headers: list[str]  # display names; blank headers become "(Column K)"
    raw_headers: list  # header cell values as they are in the file
    rows: list[list]  # data rows, padded to len(headers); fully blank rows skipped
    row_numbers: list[int]  # Excel row number of each data row
    header_row: int

    def column_index(self, header: str) -> int | None:
        wanted = header_key(header)
        for i, h in enumerate(self.headers):
            if header_key(h) == wanted:
                return i
        return None

    def column_is_empty(self, index: int) -> bool:
        return all(_is_blank(r[index]) or _is_error(r[index]) for r in self.rows)

    def to_frame(self, limit: int | None = None) -> pd.DataFrame:
        rows = self.rows if limit is None else self.rows[:limit]
        return pd.DataFrame([[display(v) for v in r] for r in rows], columns=self.headers)


class WorkbookError(ValueError):
    """A file that can't be opened as an .xlsx / .xlsm workbook (message is safe to show to the user)."""


def _open_workbook(data: bytes, **kwargs):
    if not data:
        raise WorkbookError("The file is empty (0 bytes). Upload it again.")
    if data[:8] == _OLE_SIGNATURE:
        raise WorkbookError(
            "This file is password-protected or is an old .xls file renamed to .xlsx. "
            "Open it in Excel, remove the password if it has one, and use File > Save As > Excel Workbook (.xlsx)."
        )
    try:
        return openpyxl.load_workbook(io.BytesIO(data), **kwargs)
    except (zipfile.BadZipFile, InvalidFileException) as exc:
        raise WorkbookError(
            "This is not a valid Excel workbook (it may be a CSV or another file renamed to .xlsx, or damaged). "
            "Open it in Excel and use File > Save As > Excel Workbook (.xlsx)."
        ) from exc
    except (KeyError, OSError, EOFError, zipfile.LargeZipFile) as exc:
        raise WorkbookError(f"The workbook looks damaged and can't be read ({type(exc).__name__}). "
                            "Open it in Excel and save it again.") from exc


def workbook_warnings(data: bytes) -> list[str]:
    """Things in the workbook that openpyxl cannot keep when it saves the new file."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
    except zipfile.BadZipFile:
        return []
    found = []
    if any(n.startswith(("xl/charts/", "xl/chartsheets/")) for n in names):
        found.append("charts")
    if any(n.startswith("xl/media/") for n in names) or any(
        n.startswith("xl/drawings/drawing") for n in names
    ):
        found.append("pictures / shapes")
    if any(n.startswith(("xl/slicers/", "xl/slicerCaches/", "xl/timelines/")) for n in names):
        found.append("slicers / timelines")
    return found


def read_workbook(data: bytes, header_row: int = 1) -> dict[str, Sheet]:
    wb = _open_workbook(data, read_only=True, data_only=True)
    try:
        sheets = _read_sheets(wb, header_row)
    except WorkbookError:
        raise
    except Exception as exc:  # broken XML inside the file, unsupported content, ...
        raise WorkbookError(f"The workbook looks damaged and can't be read ({type(exc).__name__}). "
                            "Open it in Excel and save it again.") from exc
    finally:
        wb.close()
    if not sheets:
        raise WorkbookError("The workbook has no worksheets.")
    return sheets


def _read_sheets(wb, header_row: int) -> dict[str, Sheet]:
    sheets: dict[str, Sheet] = {}
    for ws in wb.worksheets:
        ws.reset_dimensions()  # don't trust the stored sheet size
        raw = [list(r) for r in ws.iter_rows(min_row=header_row, values_only=True)]
        width = max((len(r) for r in raw), default=0)
        while width and all(_is_blank(r[width - 1]) for r in raw if len(r) >= width):
            width -= 1
        if not raw or not width:
            sheets[ws.title] = Sheet(ws.title, [], [], [], [], header_row)
            continue

        raw_headers = (raw[0] + [None] * width)[:width]
        headers, seen = [], set()
        for i, h in enumerate(raw_headers):
            text = re.sub(r"\s+", " ", str(h)).strip() if h is not None else ""
            text = text or f"(Column {get_column_letter(i + 1)})"
            base, n = text, 2
            while header_key(text) in seen:
                text, n = f"{base} ({n})", n + 1
            seen.add(header_key(text))
            headers.append(text)

        rows, numbers = [], []
        for offset, r in enumerate(raw[1:], start=1):
            values = (r + [None] * width)[:width]
            if all(_is_blank(v) for v in values):
                continue
            rows.append(values)
            numbers.append(header_row + offset)
        sheets[ws.title] = Sheet(ws.title, headers, raw_headers, rows, numbers, header_row)
    return sheets


# --------------------------------------------------------------------------- normalising


@dataclass
class Options:
    ignore_case: bool = True
    trim_spaces: bool = True
    smart_dates: bool = True
    duplicate_rule: str = "first"
    not_found_text: str = ""
    overwrite: bool = False
    highlight: bool = True
    add_match_column: bool = False
    add_summary_sheet: bool = False
    count_ignore_blank: bool = True  # new / not-found rows don't fail the count check


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _is_error(value) -> bool:
    """Excel error values such as #N/A (cached formula results come back as text)."""
    return isinstance(value, str) and value.strip().upper() in EXCEL_ERRORS


def _clean(value):
    return None if _is_error(value) else value


def display(value) -> str:
    """Human-friendly text for previews and joined values."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%m/%d/%Y") if value.time() == dt.time(0) else value.strftime("%m/%d/%Y %H:%M")
    if isinstance(value, dt.date):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def normalize(value, opts: Options) -> str:
    """Turn a cell value into a comparable key part (dates, numbers and text treated consistently)."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value)
    if opts.trim_spaces:
        text = re.sub(r"\s+", " ", text).strip()
    if opts.smart_dates:
        text = _text_date(text) or text
    if opts.ignore_case:
        text = text.upper()
    return text


def _text_date(text: str) -> str | None:
    try:
        if m := _SLASH_DATE.match(text):
            month, day, year = (int(g) for g in m.groups())
            if year < 100:
                year += 2000 if year < 50 else 1900
            return dt.date(year, month, day).isoformat()
        if m := _ISO_DATE.match(text):
            return dt.date(*(int(g) for g in m.groups())).isoformat()
    except ValueError:
        return None
    return None


def _row_key(row: list, columns: list[int], opts: Options) -> tuple | None:
    key = tuple(normalize(row[c], opts) for c in columns)
    return None if all(part == "" for part in key) else key


# --------------------------------------------------------------------------- suggestions


def suggest_lookup_sheet(target_name: str, lookup_names: list[str]) -> str | None:
    def simple(name):
        name = re.sub(r"^(measure|msr|sheet)[\s_\-]*", "", name.strip().lower())
        return re.sub(r"[^a-z0-9]", "", name)

    def digits(name):
        found = re.findall(r"\d+", name)
        return found[-1].lstrip("0") or "0" if found else None

    for test in (
        lambda n: n == target_name,
        lambda n: n.strip().lower() == target_name.strip().lower(),
        lambda n: simple(n) == simple(target_name),
        lambda n: digits(n) is not None and digits(n) == digits(target_name),
    ):
        hits = [n for n in lookup_names if test(n)]
        if len(hits) == 1:
            return hits[0]
    return None


def suggest_keys(target: Sheet, lookup: Sheet) -> list[str]:
    common = [h for h in target.headers if lookup.column_index(h) is not None]
    preferred = [h for h in common if header_key(h) in ("PATIENT NAME", "DOB")]
    if preferred:
        return preferred
    ids = [h for h in common if re.search(r"\b(MRN|ID|ACCOUNT|ACCT)\b", header_key(h))]
    return ids[:1] or common[:1]


def suggest_return_columns(target: Sheet, lookup: Sheet) -> list[str]:
    """The trailing lookup columns the target is missing (or has but left empty)."""
    picked = []
    for header in reversed(lookup.headers):
        ti = target.column_index(header)
        if ti is not None and not target.column_is_empty(ti):
            break
        picked.append(header)
    picked.reverse()
    return picked if len(picked) < len(lookup.headers) else []


# --------------------------------------------------------------------------- matching


@dataclass
class SheetJob:
    target_sheet: str
    lookup_sheet: str
    keys: list[tuple[str, str]]  # (target column, lookup column)
    return_columns: list[str]  # lookup columns to bring back
    fallback_keys: list[tuple[str, str]] = field(default_factory=list)
    unique_key: str | None = None  # header of a CONCATENATE-style key column to add to both files
    unique_separator: str = " | "
    unique_columns: list[tuple[str, str]] = field(default_factory=list)  # (target, lookup) joined in order; empty = keys


@dataclass
class DestColumn:
    header: str  # lookup column it is filled from
    index: int  # 0-based column in the target sheet
    is_new: bool
    raw_header: object  # header value written when the column is new


@dataclass
class CountCheck:
    column: str  # lookup column that was filled
    rows: list[dict]  # Value / Lookup file / Filled file / Difference
    ignore_blank: bool = False

    def is_ignored(self, row: dict) -> bool:
        return self.ignore_blank and row["Value"] == BLANK_LABEL

    @property
    def ok(self) -> bool:
        return all(r["Difference"] == 0 for r in self.rows if not self.is_ignored(r))


@dataclass
class SheetResult:
    job: SheetJob
    header_row: int
    row_numbers: list[int]
    dest: list[DestColumn]
    values: list[list]  # per target row, per dest column: value to write or _KEEP
    statuses: list[str]
    final: list[list]  # per target row, per dest column: value the cell ends up with
    counts: dict[str, int]
    cells_written: int
    cells_kept: int
    lookup_duplicate_keys: int
    conflicting_keys: int
    target_frame: pd.DataFrame
    lookup_rows: int
    lookup_unused: int  # lookup rows no target row matched
    count_checks: list[CountCheck]
    ignore_blank: bool
    target_duplicate_keys: int = 0  # keys that appear more than once in the sheet to fill
    unique_dest: DestColumn | None = None  # UNIQUE KEY column in the sheet to fill
    unique_values: list[str] = field(default_factory=list)
    lookup_unique_dest: DestColumn | None = None  # UNIQUE KEY column in the lookup sheet
    lookup_unique_values: list[str] = field(default_factory=list)
    lookup_row_numbers: list[int] = field(default_factory=list)
    lookup_header_row: int = 1

    @property
    def total(self) -> int:
        return len(self.statuses)

    @property
    def counts_match(self) -> bool:
        rows_ok = self.lookup_unused == 0 if self.ignore_blank else self.lookup_rows == self.total
        return rows_ok and all(c.ok for c in self.count_checks)

    def count_problems(self) -> list[str]:
        """Plain-language reasons why the counts differ between the two files."""
        problems = []
        if not self.ignore_blank:
            if self.lookup_rows != self.total:
                problems.append(f"lookup sheet has {self.lookup_rows:,} rows, filled sheet has {self.total:,}")
            if self.counts[NOT_FOUND]:
                problems.append(f"{self.counts[NOT_FOUND]:,} rows not found")
        if self.lookup_unused:
            problems.append(f"{self.lookup_unused:,} lookup rows were not used")
        for check in self.count_checks:
            diffs = [f"{r['Value']} {r['Difference']:+,}" for r in check.rows if r["Difference"] and not check.is_ignored(r)]
            if diffs:
                problems.append(f"{check.column}: " + ", ".join(diffs))
        return problems

    @property
    def matched(self) -> int:
        return self.counts[MATCHED] + self.counts[FALLBACK]

    @property
    def match_rate(self) -> float:
        considered = self.total - self.counts[BLANK_KEY]
        return self.matched / considered if considered else 0.0

    def preview(self) -> pd.DataFrame:
        frame = self.target_frame.copy()
        for i, d in enumerate(self.dest):
            label = d.header if d.is_new else frame.columns[d.index]
            frame[label] = [display(r[i]) for r in self.final]
        if self.unique_dest is not None:
            frame[self.unique_dest.header] = self.unique_values
        frame[MATCH_COLUMN] = self.statuses
        return frame


def _build_index(lookup: Sheet, columns: list[int], opts: Options) -> dict[tuple, list[list]]:
    index: dict[tuple, list[list]] = {}
    for row in lookup.rows:
        key = _row_key(row, columns, opts)
        if key is not None:
            index.setdefault(key, []).append(row)
    return index


def _pick(rows: list[list], column: int, rule: str):
    values = [_clean(r[column]) for r in rows]
    if rule == "first":
        return values[0]
    if rule == "last":
        return values[-1]
    non_blank = [v for v in values if not _is_blank(v)]
    if rule == "first_nonblank":
        return non_blank[0] if non_blank else None
    unique: list[str] = []
    for v in non_blank:
        text = display(v).strip()
        if text not in unique:
            unique.append(text)
    return " | ".join(unique) if unique else None


def _resolve_columns(sheet: Sheet, headers: list[str]) -> list[int]:
    indexes = [sheet.column_index(h) for h in headers]
    missing = [h for h, i in zip(headers, indexes) if i is None]
    if missing:
        raise ValueError(f"Sheet '{sheet.name}' has no column named: {', '.join(missing)}")
    return indexes


def run_job(target: Sheet, lookup: Sheet, job: SheetJob, opts: Options) -> SheetResult:
    if not job.keys:
        raise ValueError(f"Choose at least one key column for '{job.target_sheet}'.")
    if not job.return_columns:
        raise ValueError(f"Choose at least one column to bring back for '{job.target_sheet}'.")

    if job.unique_key:
        taken = [c for pair in job.keys + job.fallback_keys + job.unique_columns for c in pair] + job.return_columns
        if header_key(job.unique_key) in {header_key(c) for c in taken}:
            raise ValueError(
                f"'{job.target_sheet}': the UNIQUE KEY column can't be named '{job.unique_key}' because that column "
                "is used as a key or a column to fill. Give it another name."
            )

    t_keys = _resolve_columns(target, [t for t, _ in job.keys])
    l_keys = _resolve_columns(lookup, [l for _, l in job.keys])
    l_return = _resolve_columns(lookup, job.return_columns)
    index = _build_index(lookup, l_keys, opts)

    fallback_index, t_fb = None, []
    if job.fallback_keys:
        t_fb = _resolve_columns(target, [t for t, _ in job.fallback_keys])
        fallback_index = _build_index(lookup, _resolve_columns(lookup, [l for _, l in job.fallback_keys]), opts)

    dest, next_new = [], len(target.headers)
    for header, li in zip(job.return_columns, l_return):
        ti = target.column_index(header)
        if ti is None:
            raw = lookup.raw_headers[li] if not _is_blank(lookup.raw_headers[li]) else None
            dest.append(DestColumn(header, next_new, True, raw))
            next_new += 1
        else:
            dest.append(DestColumn(header, ti, False, target.raw_headers[ti]))

    counts = {MATCHED: 0, FALLBACK: 0, NOT_FOUND: 0, BLANK_KEY: 0}
    values, statuses, final = [], [], []
    written = kept = 0
    used_rows: set[int] = set()
    for row in target.rows:
        key = _row_key(row, t_keys, opts)
        found = None
        if key is None:
            status = BLANK_KEY
        elif key in index:
            status, found = MATCHED, index[key]
        elif fallback_index is not None and (fb := _row_key(row, t_fb, opts)) in fallback_index:
            status, found = FALLBACK, fallback_index[fb]
        else:
            status = NOT_FOUND
        counts[status] += 1
        statuses.append(status)

        if found is not None:
            used_rows.update(id(r) for r in found)

        row_values, row_final = [], []
        for d, li in zip(dest, l_return):
            existing = row[d.index] if d.index < len(row) else None
            if status == BLANK_KEY:
                new = None if _is_error(existing) else _KEEP  # still clear a leftover #N/A
            elif found is not None:
                new = _pick(found, li, opts.duplicate_rule)
            else:
                new = opts.not_found_text or None
            if new is not _KEEP and not opts.overwrite and not _is_blank(existing) and not _is_error(existing):
                new = _KEEP
                kept += 1
            elif new is not _KEEP and new is not None:
                written += 1
            row_values.append(new)
            row_final.append(existing if new is _KEEP else new)
        values.append(row_values)
        final.append(row_final)

    duplicate_keys = [rows for rows in index.values() if len(rows) > 1]
    conflicting = sum(
        1 for rows in duplicate_keys if any(len({display(_clean(r[li])).strip() for r in rows}) > 1 for li in l_return)
    )
    count_checks = [
        _count_check(header, [r[li] for r in lookup.rows], [r[i] for r in final], opts.count_ignore_blank)
        for i, (header, li) in enumerate(zip(job.return_columns, l_return))
    ]
    target_keys = Counter(k for k in (_row_key(r, t_keys, opts) for r in target.rows) if k is not None)

    unique = {}
    if job.unique_key:
        def concat(sheet: Sheet, columns: list[int]) -> list[str]:
            return [job.unique_separator.join(display(r[c]).strip() for c in columns) for r in sheet.rows]

        def unique_column(sheet: Sheet, next_free: int) -> DestColumn:
            existing = sheet.column_index(job.unique_key)
            if existing is not None:
                return DestColumn(job.unique_key, existing, False, sheet.raw_headers[existing])
            return DestColumn(job.unique_key, next_free, True, job.unique_key)

        t_join, l_join = t_keys, l_keys
        if job.unique_columns:
            t_join = _resolve_columns(target, [t for t, _ in job.unique_columns])
            l_join = _resolve_columns(lookup, [l for _, l in job.unique_columns])
        unique = dict(
            unique_dest=unique_column(target, next_new),
            unique_values=concat(target, t_join),
            lookup_unique_dest=unique_column(lookup, len(lookup.headers)),
            lookup_unique_values=concat(lookup, l_join),
        )
    return SheetResult(
        job=job,
        header_row=target.header_row,
        row_numbers=target.row_numbers,
        dest=dest,
        values=values,
        statuses=statuses,
        final=final,
        counts=counts,
        cells_written=written,
        cells_kept=kept,
        lookup_duplicate_keys=len(duplicate_keys),
        conflicting_keys=conflicting,
        target_frame=target.to_frame(),
        lookup_rows=len(lookup.rows),
        lookup_unused=sum(1 for r in lookup.rows if id(r) not in used_rows),
        count_checks=count_checks,
        ignore_blank=opts.count_ignore_blank,
        target_duplicate_keys=sum(1 for n in target_keys.values() if n > 1),
        lookup_row_numbers=lookup.row_numbers,
        lookup_header_row=lookup.header_row,
        **unique,
    )


def _value_counts(values) -> Counter:
    return Counter(display(_clean(v)).strip() or BLANK_LABEL for v in values)


def _count_check(column: str, lookup_values: list, filled_values: list, ignore_blank: bool) -> CountCheck:
    """Compare how often each value appears in the lookup column vs the filled column."""
    lc, fc = _value_counts(lookup_values), _value_counts(filled_values)
    order = sorted(set(lc) | set(fc), key=lambda v: (v == BLANK_LABEL, -lc.get(v, 0), v))
    rows = [
        {"Value": v, "Lookup file": lc.get(v, 0), "Filled file": fc.get(v, 0), "Difference": fc.get(v, 0) - lc.get(v, 0)}
        for v in order
    ]
    return CountCheck(column, rows, ignore_blank)


# --------------------------------------------------------------------------- writing


def _copy_style(src, dst) -> None:
    dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
    dst.alignment, dst.number_format = copy(src.alignment), src.number_format


def _safe_value(value):
    """Strip characters Excel can't store and cut text to Excel's 32,767-character cell limit."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > EXCEL_MAX_TEXT:
            value = value[: EXCEL_MAX_TEXT - 3] + "..."
    return value


def _set(ws, row: int, col: int, value):
    """Write one cell. A merged block covering it is unmerged first (merged cells can't be written)."""
    cell = ws.cell(row, col)
    if isinstance(cell, MergedCell):
        for merged in list(ws.merged_cells.ranges):
            if cell.coordinate in merged:
                ws.unmerge_cells(str(merged))
                break
        cell = ws.cell(row, col)
    cell.value = _safe_value(value)
    return cell


def _header_writer(ws, hr: int):
    """Returns add_header(col, value), styling new headers like the sheet's last existing header."""
    last_header = max((c.column for c in ws[hr] if not _is_blank(c.value)), default=1)
    template = ws.cell(hr, last_header)
    template_width = ws.column_dimensions[get_column_letter(last_header)].width

    def add_header(col: int, value) -> None:
        cell = _set(ws, hr, col, value)
        _copy_style(template, cell)
        if template_width:
            ws.column_dimensions[get_column_letter(col)].width = template_width

    return add_header


def _write_unique(ws, add_header, dest: DestColumn, row_numbers: list[int], values: list[str]) -> None:
    if dest.is_new:
        add_header(dest.index + 1, dest.raw_header)
    for row_num, value in zip(row_numbers, values):
        _set(ws, row_num, dest.index + 1, value)


def write_lookup_keys(wb, results: list[SheetResult]) -> None:
    """Add each job's UNIQUE KEY column to its lookup sheet (workbook already opened)."""
    for res in results:
        if res.lookup_unique_dest is None or res.job.lookup_sheet not in wb.sheetnames:
            continue
        ws = wb[res.job.lookup_sheet]
        add_header = _header_writer(ws, res.lookup_header_row)
        _write_unique(ws, add_header, res.lookup_unique_dest, res.lookup_row_numbers, res.lookup_unique_values)


def _save(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_lookup_output(lookup_bytes: bytes, results: list[SheetResult], keep_vba: bool = False) -> bytes:
    wb = _open_workbook(lookup_bytes, keep_vba=keep_vba)
    write_lookup_keys(wb, results)
    return _save(wb)


def build_output(
    target_bytes: bytes, results: list[SheetResult], opts: Options, keep_vba: bool = False,
    lookup_in_same_file: bool = False,
) -> bytes:
    wb = _open_workbook(target_bytes, keep_vba=keep_vba)
    for res in results:
        if res.job.target_sheet not in wb.sheetnames:
            raise ValueError(f"Sheet '{res.job.target_sheet}' is not in the file to fill. Upload the file again.")
        ws = wb[res.job.target_sheet]
        hr = res.header_row
        add_header = _header_writer(ws, hr)

        for d in res.dest:
            if d.is_new:
                add_header(d.index + 1, d.raw_header)
        if res.unique_dest is not None:
            _write_unique(ws, add_header, res.unique_dest, res.row_numbers, res.unique_values)

        match_col = None
        if opts.add_match_column:
            existing = [c.column for c in ws[hr] if not _is_blank(c.value) and header_key(c.value) == MATCH_COLUMN]
            used = [d.index + 1 for d in res.dest] + ([res.unique_dest.index + 1] if res.unique_dest else [])
            match_col = existing[0] if existing else max([ws.max_column] + used) + 1
            add_header(match_col, MATCH_COLUMN)

        for row_num, row_values, status in zip(res.row_numbers, res.values, res.statuses):
            fill = {NOT_FOUND: NOT_FOUND_FILL, FALLBACK: FALLBACK_FILL}.get(status) if opts.highlight else None
            for d, value in zip(res.dest, row_values):
                if value is _KEEP:
                    continue
                cell = _set(ws, row_num, d.index + 1, value)
                if isinstance(value, (dt.date, dt.datetime)):
                    cell.number_format = "mm/dd/yyyy"
                if fill is not None:
                    cell.fill = copy(fill)
            if match_col:
                _set(ws, row_num, match_col, status)

    if lookup_in_same_file:
        write_lookup_keys(wb, results)
    if opts.add_summary_sheet:
        _write_summary(wb, results)
    return _save(wb)


def _write_summary(wb, results: list[SheetResult]) -> None:
    if SUMMARY_SHEET in wb.sheetnames:
        del wb[SUMMARY_SHEET]
    ws = wb.create_sheet(SUMMARY_SHEET)
    ws.append([f"VLOOKUP run {dt.datetime.now():%m/%d/%Y %H:%M}"])
    ws.append([])
    headers = ["Sheet filled", "Lookup sheet", "Key columns", "Columns filled", "Rows",
               "Matched", "Fallback", "Not found", "Blank key", "Match rate", "Count check"]
    ws.append(headers)
    for res in results:
        ws.append([
            res.job.target_sheet,
            res.job.lookup_sheet,
            " + ".join(t if t == l else f"{t} = {l}" for t, l in res.job.keys),
            ", ".join(res.job.return_columns),
            res.total,
            res.counts[MATCHED],
            res.counts[FALLBACK],
            res.counts[NOT_FOUND],
            res.counts[BLANK_KEY],
            res.match_rate,
            "Same" if res.counts_match else "DIFFERENT: " + "; ".join(res.count_problems()),
        ])
        ws.cell(ws.max_row, len(headers) - 1).number_format = "0.0%"
        if not res.counts_match:
            ws.cell(ws.max_row, len(headers)).fill = copy(NOT_FOUND_FILL)
    ws["A1"].font = Font(bold=True, size=13)
    header_fill = PatternFill("solid", fgColor="1E3A8A")
    for cell in ws[3]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for i, width in enumerate([16, 16, 28, 28, 8, 10, 10, 10, 10, 11, 60], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
