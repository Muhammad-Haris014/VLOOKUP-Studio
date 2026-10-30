"""VLOOKUP Studio: fill the last columns of many Excel sheets from another workbook."""
from __future__ import annotations

import hashlib
import hmac
import html
import logging
import os
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.utils import get_column_letter

import vlookup_engine as ve

st.set_page_config(page_title="VLOOKUP Studio", page_icon=":material/join_inner:", layout="wide")

log = logging.getLogger("vlookup_studio")
STYLE_CELL_LIMIT = 200_000  # bigger tables are shown without colours (Streamlit refuses to style more)

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700&display=swap');
html, body, p, h1, h2, h3, h4, h5, label, input, textarea, .hero, .step, .kpi, .how, .chip { font-family: 'Inter', sans-serif; }
[data-testid="stAppDeployButton"], footer { display: none; }
.block-container { padding-top: 1.6rem; padding-bottom: 4rem; max-width: 1320px; }

.hero {
  background: linear-gradient(120deg, #0F172A 0%, #1E3A8A 45%, #2563EB 80%, #0EA5E9 100%);
  border-radius: 20px; padding: 30px 34px; color: #fff; margin-bottom: 8px;
  box-shadow: 0 12px 32px -12px rgba(30, 58, 138, .45);
  position: relative; overflow: hidden;
}
.hero:after {
  content: ""; position: absolute; right: -60px; top: -60px; width: 260px; height: 260px;
  border-radius: 50%; background: rgba(255,255,255,.08);
}
.hero h1 { color: #fff; font-size: 1.9rem; font-weight: 700; margin: 0; padding: 0; letter-spacing: -.02em; }
.hero p { color: rgba(255,255,255,.86); margin: 6px 0 0; font-size: 1rem; }
.hero .badges { margin-top: 16px; }
.hero .badge {
  display: inline-block; padding: 4px 12px; margin-right: 8px; border-radius: 999px;
  background: rgba(255,255,255,.14); border: 1px solid rgba(255,255,255,.22);
  font-size: .78rem; font-weight: 500;
}

.step { display: flex; align-items: center; gap: 14px; margin: 30px 0 12px; }
.step-num {
  flex: 0 0 34px; height: 34px; border-radius: 10px; background: #2563EB; color: #fff;
  display: flex; align-items: center; justify-content: center; font-weight: 700;
  box-shadow: 0 4px 12px -4px rgba(37, 99, 235, .6);
}
.step-title { font-size: 1.12rem; font-weight: 650; color: #0F172A; line-height: 1.2; }
.step-sub { color: #64748B; font-size: .88rem; margin-top: 2px; }

.chip {
  display: inline-block; padding: 3px 10px; margin: 0 6px 6px 0; border-radius: 999px;
  background: #EFF6FF; color: #1D4ED8; font-size: .76rem; font-weight: 600; border: 1px solid #DBEAFE;
}
.chip.gray { background: #F1F5F9; color: #334155; border-color: #E2E8F0; }
.chip.green { background: #ECFDF5; color: #047857; border-color: #D1FAE5; }
.chip.amber { background: #FFFBEB; color: #B45309; border-color: #FDE68A; }

.kpi {
  background: #fff; border: 1px solid #E2E8F0; border-radius: 16px; padding: 18px 20px;
  box-shadow: 0 1px 2px rgba(15, 23, 42, .04); border-top: 4px solid #2563EB; min-height: 128px;
}
.kpi.good { border-top-color: #10B981; }
.kpi.warn { border-top-color: #F59E0B; }
.kpi.info { border-top-color: #6366F1; }
.kpi-label { color: #64748B; font-size: .74rem; font-weight: 600; text-transform: uppercase; letter-spacing: .06em; }
.kpi-value { color: #0F172A; font-size: clamp(1.3rem, 2.4vw, 1.85rem); white-space: nowrap; font-weight: 700; margin-top: 4px; letter-spacing: -.02em; }
.kpi-hint { color: #94A3B8; font-size: .8rem; margin-top: 2px; }

.how { background: #fff; border: 1px solid #E2E8F0; border-radius: 16px; padding: 20px; height: 100%; }
.how .icon { font-size: 1.5rem; }
.how h4 { margin: 8px 0 4px; font-size: 1rem; color: #0F172A; }
.how p { color: #64748B; font-size: .88rem; margin: 0; }

.dest { font-size: .85rem; color: #334155; margin: 2px 0; }
.dest code { background: #F1F5F9; padding: 1px 6px; border-radius: 6px; }

[data-testid="stSidebar"] { background: #0F172A; }
[data-testid="stSidebar"] * { color: #E2E8F0; }
[data-testid="stSidebar"] [data-baseweb="select"] *, [data-testid="stSidebar"] input { color: #0F172A !important; }
[data-testid="stSidebar"] .stCaption, [data-testid="stSidebar"] small { color: #94A3B8 !important; }
[data-testid="stSidebar"] hr { border-color: #1E293B; }
[data-testid="stFileUploaderDropzone"] { border-radius: 14px; }
div[data-testid="stVerticalBlockBorderWrapper"] { border-radius: 16px; }
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

STATUS_COLORS = {
    ve.MATCHED: "background-color: #ECFDF5; color: #047857",
    ve.FALLBACK: "background-color: #FFF7ED; color: #C2410C",
    ve.NOT_FOUND: "background-color: #FEF3C7; color: #92400E",
    ve.BLANK_KEY: "background-color: #F1F5F9; color: #475569",
}


def step(num: int, title: str, subtitle: str = "") -> None:
    st.markdown(
        f'<div class="step"><div class="step-num">{num}</div><div>'
        f'<div class="step-title">{title}</div><div class="step-sub">{subtitle}</div></div></div>',
        unsafe_allow_html=True,
    )


def kpi(col, label: str, value: str, hint: str = "", tone: str = "") -> None:
    col.markdown(
        f'<div class="kpi {tone}"><div class="kpi-label">{label}</div>'
        f'<div class="kpi-value">{value}</div><div class="kpi-hint">{hint}</div></div>',
        unsafe_allow_html=True,
    )


def chips(items, tone: str = "") -> str:
    return "".join(f'<span class="chip {tone}">{item}</span>' for item in items)


@st.cache_data(show_spinner=False, max_entries=6, ttl=3600)
def load(data: bytes, header_row: int) -> dict[str, ve.Sheet]:
    return ve.read_workbook(data, header_row)


def required_password() -> str:
    """Optional sign-in password: `password` in .streamlit/secrets.toml or the VLOOKUP_PASSWORD variable."""
    if st.secrets.load_if_toml_exists() and "password" in st.secrets:
        return str(st.secrets["password"])
    return os.environ.get("VLOOKUP_PASSWORD", "")


def sign_in_gate() -> None:
    password = required_password()
    if not password or st.session_state.get("signed_in"):
        return
    _, middle, _ = st.columns([1, 1.2, 1])
    with middle, st.form("sign_in"):
        st.markdown("### :material/lock: VLOOKUP Studio")
        typed = st.text_input("Password", type="password")
        if st.form_submit_button("Sign in", type="primary", width="stretch"):
            if hmac.compare_digest(typed.encode(), password.encode()):
                st.session_state["signed_in"] = True
                st.rerun()
            st.error("Wrong password.")
    st.stop()


def show_error(action: str, exc: Exception) -> None:
    """Friendly message for expected problems; full details only in the server log and an expander."""
    if isinstance(exc, ValueError):  # includes WorkbookError: messages written for the user
        st.error(f"{action}: {exc}", icon=":material/error:")
        return
    log.exception("%s failed", action)
    st.error(f"{action}: something unexpected went wrong ({type(exc).__name__}). "
             "Nothing was changed. Try again, or send the details below to whoever looks after this tool.",
             icon=":material/error:")
    with st.expander("Technical details"):
        st.exception(exc)


def styled(frame: pd.DataFrame, style):
    """Apply a Styler only when the table is small enough for Streamlit to render it."""
    return style(frame.style) if frame.size <= STYLE_CELL_LIMIT else frame


def sheet_plan_chips(sheet_names: list[str], jobs: list[ve.SheetJob]) -> str:
    filled = {j.target_sheet: j.lookup_sheet for j in jobs}
    return "".join(
        chips([f"{html.escape(n)} ← {html.escape(filled[n])}"], "green") if n in filled
        else chips([f"{html.escape(n)} · kept as is"], "gray")
        for n in sheet_names
    )


def same_name_pairs(lookup: ve.Sheet, columns: list[str]) -> list[tuple[str, str]]:
    """(target column, lookup column) pairs for columns that have the same name in the lookup sheet."""
    return [(c, lookup.headers[lookup.column_index(c)]) for c in columns]


def excel_mime(suffix: str) -> str:
    if suffix.lower() == ".xlsm":
        return "application/vnd.ms-excel.sheet.macroEnabled.12"
    return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def digest(data: bytes) -> str:
    return hashlib.md5(data).hexdigest()[:12]


def apply_to_all(source: str, others: list[tuple[str, list[str], list[str]]]) -> None:
    """Copy one sheet's key / return settings to every other sheet where the columns exist."""
    ss = st.session_state
    keys, ret, fb = ss[f"{source}::keys"], ss[f"{source}::ret"], ss[f"{source}::fb"]
    for prefix, target_headers, lookup_headers in others:
        t_keys = {ve.header_key(h): h for h in target_headers}
        l_keys = {ve.header_key(h): h for h in lookup_headers}
        ss[f"{prefix}::keys"] = [t_keys[ve.header_key(k)] for k in keys if ve.header_key(k) in t_keys]
        ss[f"{prefix}::fb"] = [t_keys[ve.header_key(k)] for k in fb if ve.header_key(k) in t_keys]
        ss[f"{prefix}::ret"] = [l_keys[ve.header_key(c)] for c in ret if ve.header_key(c) in l_keys]
        for setting in ("uk", "ukname", "uksep", "ukmatch"):
            ss[f"{prefix}::{setting}"] = ss[f"{source}::{setting}"]
        ss[f"{prefix}::ukcols"] = [
            t_keys[ve.header_key(c)] for c in ss[f"{source}::ukcols"]
            if ve.header_key(c) in t_keys and ve.header_key(c) in l_keys
        ]
        for k in keys:
            chosen = ss.get(f"{source}::map::{k}")
            if ve.header_key(k) in t_keys and chosen and ve.header_key(chosen) in l_keys:
                ss[f"{prefix}::map::{t_keys[ve.header_key(k)]}"] = l_keys[ve.header_key(chosen)]


sign_in_gate()

# --------------------------------------------------------------------------- sidebar

with st.sidebar:
    st.markdown("### :material/tune: Settings")
    st.markdown("**Header row**")
    target_header = st.number_input("File to fill", min_value=1, max_value=50, value=1)
    lookup_header = st.number_input("Lookup file", min_value=1, max_value=50, value=1)
    st.divider()
    st.markdown("**Matching**")
    ignore_case = st.toggle("Ignore upper / lower case", value=True)
    trim_spaces = st.toggle("Ignore extra spaces", value=True)
    smart_dates = st.toggle(
        "Match dates in any format", value=True,
        help="01/20/2026, 2026-01-20 and a real Excel date are treated as the same value.",
    )
    duplicate_rule = st.selectbox(
        "If a key is found more than once",
        options=list(ve.DUPLICATE_RULES), format_func=ve.DUPLICATE_RULES.get,
    )
    st.divider()
    st.markdown("**Output**")
    not_found_text = st.text_input(
        "Write this when not found", value="",
        help="Empty (default) leaves the cell blank, so no #N/A is written. "
             "#N/A and other Excel errors are always removed automatically.",
        placeholder="blank, no #N/A",
    )
    overwrite = st.toggle("Overwrite cells that already have a value", value=False)
    highlight = st.toggle("Highlight not-found cells", value=True)
    add_match_column = st.toggle("Add a MATCH STATUS column", value=False)
    add_summary_sheet = st.toggle("Add a summary sheet", value=False)
    st.divider()
    st.markdown("**Count check**")
    count_ignore_blank = st.toggle(
        "Ignore blank / not-found rows", value=True,
        help="New patients that are not in the lookup file stay blank and do not cause a warning. "
             "You are still warned if PM / PNM (or any other value) counts differ, or a lookup row was not used.",
    )
    st.divider()
    st.caption(":material/lock: Files are processed in memory only. Nothing is saved on the server "
               "and your original files are never changed.")

opts = ve.Options(
    ignore_case=ignore_case, trim_spaces=trim_spaces, smart_dates=smart_dates,
    duplicate_rule=duplicate_rule, not_found_text=not_found_text, overwrite=overwrite,
    highlight=highlight, add_match_column=add_match_column, add_summary_sheet=add_summary_sheet,
    count_ignore_blank=count_ignore_blank,
)

# --------------------------------------------------------------------------- header

st.markdown(
    """
    <div class="hero">
      <h1>VLOOKUP Studio</h1>
      <p>Fill STATUS, COMMENTS or any last columns across many sheets at once, matched on the columns you choose.</p>
      <div class="badges">
        <span class="badge">Multi-sheet</span><span class="badge">Multi-column keys</span>
        <span class="badge">Fallback matching</span><span class="badge">Original file untouched</span>
      </div>
    </div>
    """,
    unsafe_allow_html=True,
)

# --------------------------------------------------------------------------- step 1: files

step(1, "Upload your files", "The file to fill gets new values; the lookup file is where the values come from.")
col_t, col_l = st.columns(2, gap="large")
with col_t, st.container(border=True):
    st.markdown("**:material/edit_document: File to fill**")
    st.caption("New data file. Values are written into its last columns.")
    target_file = st.file_uploader("File to fill", type=["xlsx", "xlsm"], key="target_upload", label_visibility="collapsed")
with col_l, st.container(border=True):
    st.markdown("**:material/database: Lookup file**")
    st.caption("Old / finished file that already has STATUS, COMMENTS, etc.")
    same_file = st.checkbox("Lookup values are in the same file")
    lookup_file = None
    if not same_file:
        lookup_file = st.file_uploader("Lookup file", type=["xlsx", "xlsm"], key="lookup_upload", label_visibility="collapsed")

if target_file is None or (lookup_file is None and not same_file):
    st.write("")
    how = st.columns(4, gap="medium")
    for col, (icon, title, text) in zip(how, [
        ("📂", "1 · Upload", "Pick the file to fill and the file that has the answers."),
        ("🔗", "2 · Pair sheets", "Sheets with the same name are paired automatically, e.g. 001 ↔ 001."),
        ("🔑", "3 · Choose keys", "Match on one or more columns, e.g. PATIENT NAME + DOB."),
        ("⬇️", "4 · Download", "Get a new workbook with the last columns filled in."),
    ]):
        col.markdown(f'<div class="how"><div class="icon">{icon}</div><h4>{title}</h4><p>{text}</p></div>', unsafe_allow_html=True)
    st.stop()

target_bytes = target_file.getvalue()
lookup_bytes = target_bytes if same_file else lookup_file.getvalue()
lookup_name = target_file.name if same_file else lookup_file.name

sheets_by_file = {}
for label, name, data, header_row in (
    ("file to fill", target_file.name, target_bytes, target_header),
    ("lookup file", lookup_name, lookup_bytes, lookup_header),
):
    try:
        with st.spinner(f"Reading {name}…"):
            sheets_by_file[label] = load(data, int(header_row))
    except Exception as exc:  # unreadable / corrupt / protected workbook
        show_error(f"Could not read the {label} “{name}”", exc)
        st.stop()
t_sheets, l_sheets = sheets_by_file["file to fill"], sheets_by_file["lookup file"]

for col, name, sheets in ((col_t, target_file.name, t_sheets), (col_l, lookup_name, l_sheets)):
    rows = sum(len(s.rows) for s in sheets.values())
    col.markdown(
        chips([f"{len(sheets)} sheets", f"{rows:,} rows"], "green") + chips([html.escape(n) for n in sheets], "gray"),
        unsafe_allow_html=True,
    )
    if not rows:
        col.warning("No data rows found. Check the header row number in the sidebar.", icon=":material/warning:")

lost = ve.workbook_warnings(target_bytes)
if lost:
    st.warning(
        f"**{target_file.name}** contains {', '.join(lost)}. These may be dropped or simplified in "
        "the new workbook. Cell values, formatting and all sheets are kept; check them in Excel after downloading.",
        icon=":material/image_not_supported:",
    )

sig = f"{digest(target_bytes)}-{digest(lookup_bytes)}-{target_header}-{lookup_header}"
ss = st.session_state
if ss.get("last_run") and ss["last_run"]["sig"] != sig:
    del ss["last_run"]  # results of other files: drop them so old patient data isn't kept in memory

# --------------------------------------------------------------------------- step 2: sheet pairing

step(2, "Pair the sheets", "Tick each sheet you want to fill and pick which lookup sheet its values come from.")
lookup_names = list(l_sheets)
default_rows = []
for name, sheet in t_sheets.items():
    suggestion = ve.suggest_lookup_sheet(name, lookup_names)
    if same_file and suggestion == name:
        suggestion = None
    use = bool(suggestion) and bool(sheet.rows) and bool(ve.suggest_return_columns(sheet, l_sheets[suggestion]))
    default_rows.append({"Fill": use, "Sheet to fill": name, "Rows": len(sheet.rows), "Lookup sheet": suggestion})

with st.container(border=True):
    mapping = st.data_editor(
        pd.DataFrame(default_rows),
        key=f"mapping-{sig}",
        hide_index=True,
        width="stretch",
        disabled=["Sheet to fill", "Rows"],
        column_config={
            "Fill": st.column_config.CheckboxColumn("Fill", width="small", help="Tick to fill this sheet."),
            "Sheet to fill": st.column_config.TextColumn("Sheet to fill"),
            "Rows": st.column_config.NumberColumn("Rows", width="small"),
            "Lookup sheet": st.column_config.SelectboxColumn(
                "Get values from lookup sheet", options=lookup_names, help="Sheet in the lookup file to search."
            ),
        },
    )
    st.caption("Tip: sheets with the same name or number are paired for you. Click a cell in the last column to change it.")

pairs = []
for _, row in mapping.iterrows():
    if not row["Fill"]:
        continue
    if not row["Lookup sheet"] or pd.isna(row["Lookup sheet"]):
        st.warning(f"Sheet **{row['Sheet to fill']}** is ticked but has no lookup sheet. Pick one in the table.")
        continue
    pairs.append((row["Sheet to fill"], row["Lookup sheet"]))

if not pairs:
    st.info("Tick at least one sheet in the table above to continue.", icon=":material/info:")
    st.stop()

# --------------------------------------------------------------------------- step 3: keys and columns

step(3, "Choose the key and the columns to fill",
     "The key (primary) columns decide which rows match. The chosen columns are copied back, like VLOOKUP.")

prefixes = {pair: f"{sig}::{pair[0]}::{pair[1]}" for pair in pairs}
jobs: list[ve.SheetJob] = []
problems: list[str] = []
tabs = st.tabs([f"{t}  ←  {l}" for t, l in pairs])
for tab, (t_name, l_name) in zip(tabs, pairs):
    target, lookup = t_sheets[t_name], l_sheets[l_name]
    p = prefixes[(t_name, l_name)]
    if not target.headers or not lookup.headers:
        empty = t_name if not target.headers else f"{l_name} (lookup)"
        with tab:
            st.warning(f"Sheet **{empty}** has no header row / data, so there is nothing to match. "
                       "Check the header row number in the sidebar, or untick this sheet.", icon=":material/warning:")
        problems.append(f"**{t_name}**: skipped, sheet **{empty}** is empty. It is kept as it is in the latest file.")
        continue
    ss.setdefault(f"{p}::keys", ve.suggest_keys(target, lookup))
    ss.setdefault(f"{p}::ret", ve.suggest_return_columns(target, lookup))
    ss.setdefault(f"{p}::fb", [])
    ss.setdefault(f"{p}::uk", False)
    ss.setdefault(f"{p}::ukname", "UNIQUE KEY")
    ss.setdefault(f"{p}::uksep", " | ")
    ss.setdefault(f"{p}::ukcols", [k for k in ss[f"{p}::keys"] if lookup.column_index(k) is not None])
    ss.setdefault(f"{p}::ukmatch", False)

    with tab:
        left, right = st.columns(2, gap="large")
        with left, st.container(border=True):
            st.markdown("**:material/key: Primary key: match rows on**")
            keys = st.multiselect(
                "Key columns (in the sheet to fill)", target.headers, key=f"{p}::keys",
                help="Every chosen column must match. PATIENT NAME + DOB is safer than name alone.",
            )
            key_pairs = []
            unmatched = [k for k in keys if lookup.column_index(k) is None]
            with st.expander("Key column names in the lookup sheet", expanded=bool(unmatched)):
                for k in keys:
                    li = lookup.column_index(k)
                    ss.setdefault(f"{p}::map::{k}", lookup.headers[li] if li is not None else lookup.headers[0])
                    lk = st.selectbox(f"“{k}” is called", lookup.headers, key=f"{p}::map::{k}")
                    key_pairs.append((k, lk))
                if not keys:
                    st.caption("Choose key columns first.")

            common = [h for h in target.headers if lookup.column_index(h) is not None]
            fb = st.multiselect(
                "Fallback key (optional)", common, key=f"{p}::fb",
                help="Tried only when the primary key finds nothing, e.g. PATIENT NAME alone. Fallback matches are marked orange.",
            )

            make_unique = st.toggle(
                "Create a UNIQUE KEY column in both files (like CONCATENATE)", key=f"{p}::uk",
                help="Joins the columns you pick, in the order you pick them, into one new column. "
                     "It is added at the end of the sheet to fill AND of the lookup sheet.",
            )
            # Always rendered (just disabled) so the inputs keep showing their saved values.
            uk_cols = st.multiselect(
                "Columns to join (in this order)", common, key=f"{p}::ukcols", disabled=not make_unique,
                placeholder="e.g. PATIENT NAME, DOS",
                help="Pick the columns that make each row unique, e.g. PATIENT NAME + DOS (+ CPT). "
                     "They must exist in both sheets. The order you pick is the order they are joined.",
            )
            name_col, sep_col = st.columns([3, 2])
            typed_name = name_col.text_input("Column name", key=f"{p}::ukname", disabled=not make_unique)
            separator = sep_col.text_input(
                "Separator", key=f"{p}::uksep", disabled=not make_unique,
                help="Leave empty to join with nothing, like =A2&B2.",
            )
            match_on_unique = st.checkbox(
                "Match rows on this UNIQUE KEY (instead of the primary key above)", key=f"{p}::ukmatch",
                disabled=not make_unique,
                help="Works like VLOOKUP on a CONCATENATE column in Excel: rows match when the joined columns match.",
            )
            unique_name = None
            if make_unique:
                unique_name = typed_name.strip() or "UNIQUE KEY"
                existing_col = target.column_index(unique_name)
                if existing_col is not None and not target.column_is_empty(existing_col):
                    st.caption(f":material/warning: Column **{unique_name}** already has data in this sheet. "
                               "It will be replaced with the new key.")
                if uk_cols and target.rows:
                    key_idx = [target.column_index(k) for k in uk_cols]
                    example = separator.join(ve.display(target.rows[0][i]).strip() for i in key_idx)
                    t_col = target.column_index(unique_name)
                    l_col = lookup.column_index(unique_name)
                    t_letter = get_column_letter((t_col if t_col is not None else len(target.headers) + sum(
                        target.column_index(c) is None for c in ss[f"{p}::ret"])) + 1)
                    l_letter = get_column_letter((l_col if l_col is not None else len(lookup.headers)) + 1)
                    st.markdown(
                        f'<div class="dest">Example: <code>{html.escape(example)}</code></div>'
                        f'<div class="dest">• Sheet to fill → column <code>{t_letter}</code> &nbsp; '
                        f'• Lookup sheet → column <code>{l_letter}</code></div>',
                        unsafe_allow_html=True,
                    )
                else:
                    st.caption("Pick the columns to join.")

        with right, st.container(border=True):
            st.markdown("**:material/output: Columns to fill**")
            ret = st.multiselect(
                "Bring back these lookup columns", lookup.headers, key=f"{p}::ret",
                help="Usually the last columns of the lookup sheet, e.g. STATUS and COMMENTS.",
            )
            next_new = len(target.headers)
            for c in ret:
                ti = target.column_index(c)
                if ti is None:
                    where = f"new column <code>{get_column_letter(next_new + 1)}</code> at the end"
                    next_new += 1
                else:
                    where = f"existing column <code>{get_column_letter(ti + 1)}</code>"
                st.markdown(f'<div class="dest">• <b>{html.escape(c)}</b> → {where}</div>', unsafe_allow_html=True)
            if not ret:
                st.caption(":material/info: No columns chosen. This sheet is kept exactly as it is in the latest file.")
            st.write("")
            others = [(prefixes[o], t_sheets[o[0]].headers, l_sheets[o[1]].headers) for o in pairs if o != (t_name, l_name)]
            st.button(
                "Use these settings for all sheets", key=f"{p}::apply", icon=":material/content_copy:",
                on_click=apply_to_all, args=(p, others), disabled=not others,
            )

        with st.expander("Preview both sheets", icon=":material/table_view:"):
            a, b = st.columns(2, gap="medium")
            a.caption(f"Sheet to fill · {t_name} · {len(target.rows):,} rows")
            a.dataframe(target.to_frame(limit=8), hide_index=True, width="stretch")
            b.caption(f"Lookup sheet · {l_name} · {len(lookup.rows):,} rows")
            b.dataframe(lookup.to_frame(limit=8), hide_index=True, width="stretch")

    unique_pairs = same_name_pairs(lookup, uk_cols) if make_unique else []
    if make_unique and match_on_unique:
        key_pairs = unique_pairs  # match on the joined columns, like VLOOKUP on a CONCATENATE column
    if not ret:
        pass  # nothing to fill: the sheet is copied to the download as it is in the latest file
    elif make_unique and not uk_cols:
        problems.append(f"**{t_name}**: pick the columns to join for the UNIQUE KEY.")
    elif make_unique and ve.header_key(unique_name) in {
        ve.header_key(c) for c in [*keys, *fb, *uk_cols, *ret, *(lk for _, lk in key_pairs)]
    }:
        problems.append(f"**{t_name}**: the UNIQUE KEY can't be called **{unique_name}**, that column is used as a "
                        "key or a column to fill. Give it another name.")
    elif not key_pairs:
        problems.append(f"**{t_name}**: choose at least one key column.")
    else:
        jobs.append(ve.SheetJob(
            t_name, l_name, key_pairs, list(ret), same_name_pairs(lookup, fb),
            unique_key=unique_name, unique_separator=ss[f"{p}::uksep"], unique_columns=unique_pairs,
        ))

# --------------------------------------------------------------------------- step 4: run

step(4, "Run and download", "A new workbook is created. Your original files are not changed.")
for problem in problems:
    st.warning(problem, icon=":material/warning:")

with st.container(border=True):
    st.markdown(
        f"**:material/inventory_2: The download will contain all {len(t_sheets)} sheets of "
        f"{html.escape(target_file.name)}** &nbsp; " + sheet_plan_chips(list(t_sheets), jobs),
        unsafe_allow_html=True,
    )
    st.caption("Green = filled by VLOOKUP. Gray = kept exactly as it is in the latest file (no columns chosen or not ticked).")

fingerprint = repr((sig, jobs, opts))
run = st.button(
    f"Run VLOOKUP on {len(jobs)} sheet{'s' if len(jobs) != 1 else ''}",
    type="primary", icon=":material/play_arrow:", disabled=not jobs,
)
if run:
    try:
        with st.status("Running VLOOKUP…", expanded=True) as status:
            results = []
            for job in jobs:
                st.write(f"Matching **{job.target_sheet}** against **{job.lookup_sheet}**…")
                results.append(ve.run_job(t_sheets[job.target_sheet], l_sheets[job.lookup_sheet], job, opts))
            st.write("Writing the new workbook…")
            output = ve.build_output(
                target_bytes, results, opts, keep_vba=target_file.name.lower().endswith(".xlsm"),
                lookup_in_same_file=same_file,
            )
            lookup_output = None
            if not same_file and any(r.lookup_unique_dest for r in results):
                st.write("Adding UNIQUE KEY to the lookup file…")
                lookup_output = ve.build_lookup_output(
                    lookup_bytes, results, keep_vba=lookup_name.lower().endswith(".xlsm"),
                )
            status.update(label="Done", state="complete", expanded=False)
        if all(r.counts_match for r in results):
            st.toast("Count check passed: both files have the same counts.", icon=":material/check_circle:")
        else:
            st.toast("Warning: counts are NOT the same in both files. See the count check.", icon=":material/warning:")
        path, lpath = Path(target_file.name), Path(lookup_name)
        ss["last_run"] = {
            "sig": sig, "fingerprint": fingerprint, "results": results, "output": output,
            "file_name": f"{path.stem} - VLOOKUP{path.suffix}", "mime": excel_mime(path.suffix),
            "lookup_output": lookup_output,
            "lookup_file_name": f"{lpath.stem} - UNIQUE KEY{lpath.suffix}", "lookup_mime": excel_mime(lpath.suffix),
        }
    except Exception as exc:
        ss.pop("last_run", None)
        show_error("VLOOKUP did not finish", exc)

last = ss.get("last_run")
if last and last["sig"] == sig:
    if last["fingerprint"] != fingerprint:
        st.info("Settings changed since the last run. Run again to apply them.", icon=":material/refresh:")
    results: list[ve.SheetResult] = last["results"]

    total = sum(r.total for r in results)
    matched = sum(r.matched for r in results)
    fallback = sum(r.counts[ve.FALLBACK] for r in results)
    not_found = sum(r.counts[ve.NOT_FOUND] for r in results)
    considered = total - sum(r.counts[ve.BLANK_KEY] for r in results)
    rate = matched / considered if considered else 0

    st.write("")
    k1, k2, k3, k4 = st.columns(4, gap="medium")
    kpi(k1, "Rows checked", f"{total:,}", f"{len(results)} sheet{'s' if len(results) != 1 else ''}", "info")
    kpi(k2, "Matched", f"{matched:,}", f"{fallback:,} by fallback key" if fallback else "all by primary key", "good")
    kpi(k3, "Not found", f"{not_found:,}", "highlighted in the file" if opts.highlight else "", "warn")
    kpi(k4, "Match rate", f"{rate:.1%}", f"{sum(r.cells_written for r in results):,} cells written")

    st.write("")
    mismatched = [r for r in results if not r.counts_match]
    if mismatched:
        lines = "\n".join(f"- **{r.job.target_sheet}**: " + "; ".join(r.count_problems()) for r in mismatched)
        st.error(
            f"**Count check: the counts are NOT the same in both files** "
            f"({len(mismatched)} of {len(results)} sheets).\n\n{lines}\n\n"
            "Check the count tables under each sheet below before you use the file.",
            icon=":material/warning:",
        )
    elif opts.count_ignore_blank:
        ignored = sum(r.counts[ve.NOT_FOUND] for r in results)
        st.success(
            "**Count check passed:** every lookup row was used and each value (PM, PNM, …) has the same count "
            "in both files." + (f" {ignored:,} blank / not-found rows were ignored." if ignored else ""),
            icon=":material/check_circle:",
        )
    else:
        st.success(
            "**Count check passed:** every sheet has the same number of rows and the same count of each value "
            "(PM, PNM, …) in both files.",
            icon=":material/check_circle:",
        )
    summary = pd.DataFrame([{
        "Sheet": r.job.target_sheet,
        "From": r.job.lookup_sheet,
        "Key": " + ".join(t if ve.header_key(t) == ve.header_key(l) else f"{t}={l}" for t, l in r.job.keys),
        "Filled": ", ".join(r.job.return_columns),
        "Rows": r.total,
        "Matched": r.counts[ve.MATCHED],
        "Fallback": r.counts[ve.FALLBACK],
        "Not found": r.counts[ve.NOT_FOUND],
        "Match rate": round(r.match_rate * 100, 1),
        "Count check": "✅ Same" if r.counts_match else "⚠️ Different",
        "Kept (had value)": r.cells_kept,
        "Duplicate keys": r.lookup_duplicate_keys,
    } for r in results])
    st.dataframe(
        summary, hide_index=True, width="stretch",
        column_config={"Match rate": st.column_config.ProgressColumn("Match rate", min_value=0, max_value=100, format="%.1f%%")},
    )

    for r in results:
        if r.conflicting_keys:
            st.warning(
                f"**{r.job.target_sheet}**: {r.conflicting_keys:,} key(s) appear more than once in lookup sheet "
                f"**{r.job.lookup_sheet}** with different values. Rule used: *{ve.DUPLICATE_RULES[opts.duplicate_rule]}*. "
                "Add another key column (e.g. DOB or DOS) to make matches exact.",
                icon=":material/content_copy:",
            )

    for r in results:
        if r.target_duplicate_keys or (r.unique_dest is not None and r.lookup_duplicate_keys):
            st.info(
                f"**{r.job.target_sheet}**: key is not unique. {r.target_duplicate_keys:,} key(s) repeat in the sheet "
                f"to fill and {r.lookup_duplicate_keys:,} in the lookup sheet. Add more key columns (e.g. DOS) "
                "if every row should be unique.",
                icon=":material/fingerprint:",
            )

    if last.get("lookup_output"):
        d1, d2 = st.columns(2, gap="medium")
        d1.download_button(
            "Download filled workbook", data=last["output"], file_name=last["file_name"], mime=last["mime"],
            type="primary", icon=":material/download:", width="stretch",
        )
        d2.download_button(
            "Optional: old (lookup) file with UNIQUE KEY added", data=last["lookup_output"],
            file_name=last["lookup_file_name"], mime=last["lookup_mime"],
            icon=":material/download:", width="stretch",
            help="Only needed if you also want the UNIQUE KEY column in the old file. "
                 "Nothing else in that file is changed. Your result is the blue button.",
        )
    else:
        st.download_button(
            "Download filled workbook", data=last["output"], file_name=last["file_name"], mime=last["mime"],
            type="primary", icon=":material/download:", width="stretch",
        )
    st.markdown(
        f'<div class="dest" style="margin-top:6px">Includes all {len(t_sheets)} sheets: </div>'
        + sheet_plan_chips(list(t_sheets), [r.job for r in results]),
        unsafe_allow_html=True,
    )

    st.markdown('<div class="step-sub" style="margin:18px 0 6px">Result preview</div>', unsafe_allow_html=True)
    preview_tabs = st.tabs([r.job.target_sheet for r in results])
    for tab, r in zip(preview_tabs, results):
        with tab:
            with st.container(border=True):
                rows_ok = r.ignore_blank or r.lookup_rows == r.total
                not_found_chip = (
                    chips([f"Not found (ignored) {r.counts[ve.NOT_FOUND]:,}"], "gray")
                    if r.ignore_blank and r.counts[ve.NOT_FOUND] else ""
                )
                st.markdown(
                    f"**:material/fact_check: Count check** &nbsp; "
                    + chips([f"Lookup rows {r.lookup_rows:,}", f"Filled rows {r.total:,}"], "green" if rows_ok else "amber")
                    + chips([f"Lookup rows not used {r.lookup_unused:,}"], "amber" if r.lookup_unused else "green")
                    + not_found_chip,
                    unsafe_allow_html=True,
                )
                check_cols = st.columns(min(len(r.count_checks), 3) or 1, gap="medium")
                for i, check in enumerate(r.count_checks):
                    with check_cols[i % len(check_cols)]:
                        st.caption(f"{check.column} · {'same in both files' if check.ok else 'DIFFERENT'}")
                        table = pd.DataFrame(check.rows)
                        table["Note"] = ["ignored" if check.is_ignored(row) else "" for row in check.rows]

                        def diff_styles(df: pd.DataFrame, check=check) -> pd.DataFrame:
                            styles = pd.DataFrame("", index=df.index, columns=df.columns)
                            for idx, row in df.iterrows():
                                if check.is_ignored(row):
                                    styles.loc[idx, :] = "color: #94A3B8"
                                elif row["Difference"]:
                                    styles.loc[idx, "Difference"] = "background-color: #FEE2E2; color: #B91C1C; font-weight: 600"
                                else:
                                    styles.loc[idx, "Difference"] = "color: #047857"
                            return styles

                        st.dataframe(styled(table, lambda s: s.apply(diff_styles, axis=None)),
                                     hide_index=True, width="stretch")
            show = st.segmented_control(
                "Show", ["All", ve.MATCHED, ve.FALLBACK, ve.NOT_FOUND], default="All",
                key=f"show::{r.job.target_sheet}", label_visibility="collapsed",
            )
            frame = r.preview()
            if show and show != "All":
                frame = frame[frame[ve.MATCH_COLUMN] == show]
            st.caption(f"{len(frame):,} rows" + ("" if frame.size <= STYLE_CELL_LIMIT else " · large sheet, shown without colours"))
            st.dataframe(
                styled(frame, lambda s: s.map(lambda v: STATUS_COLORS.get(v, ""), subset=[ve.MATCH_COLUMN])),
                hide_index=True, width="stretch", height=420,
            )
